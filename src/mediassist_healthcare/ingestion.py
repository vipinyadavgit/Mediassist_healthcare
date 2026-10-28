"""Load one healthcare document and expose its extracted text and source details.

This module handles Phase 4: reading PDF, DOCX, and TXT files one at a time,
without chunking, embedding, or indexing their contents.
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from pypdf import PdfReader


@dataclass(frozen=True)
class LoadedDocument:
    """Keep extracted text together with the original file's basic details."""

    source: Path
    file_type: str
    text: str

# Open a PDF and collect readable text from its pages in document order.
def _read_pdf(path: Path) -> str:
    """Extract selectable text from every page in a PDF file."""

    reader = PdfReader(str(path))
    return "\n".join(
        text
        for page in reader.pages
        if (text := page.extract_text())
    )


# Open a Word document and collect text from its paragraphs and tables.
def _read_docx(path: Path) -> str:
    """Extract paragraph and table text from a Word document."""

    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    for table in document.tables:
        for row in table.rows:
            row_text = "\t".join(cell.text for cell in row.cells if cell.text)
            if row_text:
                parts.append(row_text)
    return "\n".join(parts)


# Validate one file and return its extracted text with source metadata.
def load_document(file_path: str | Path) -> LoadedDocument:
    """Validate and extract text from one supported document file."""

    source = Path(file_path)
    if not source.is_file():
        raise FileNotFoundError(f"Document does not exist or is not a file: {source}")

    suffix = source.suffix.lower()
    if suffix == ".pdf":
        text = _read_pdf(source)
        file_type = "pdf"
    elif suffix == ".docx":
        text = _read_docx(source)
        file_type = "docx"
    elif suffix == ".txt":
        text = source.read_text(encoding="utf-8-sig")
        file_type = "txt"
    else:
        raise ValueError(
            f"Unsupported document type '{suffix or '(no extension)'}'. "
            "Supported types are PDF, DOCX, and TXT."
        )

    return LoadedDocument(source=source, file_type=file_type, text=text)


# Provide a command-line interface for extracting text from one document.
def main() -> None:
    """Read one command-line file path and print its extracted text."""

    parser = argparse.ArgumentParser(
        description="Extract text from one PDF, DOCX, or TXT document."
    )
    parser.add_argument("file", type=Path, help="Path to one supported document")
    args = parser.parse_args()

    document = load_document(args.file)
    print(f"Source: {document.source}")
    print(f"Type: {document.file_type}")
    print()
    print(document.text)


if __name__ == "__main__":
    main()
