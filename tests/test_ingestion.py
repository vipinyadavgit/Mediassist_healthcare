"""Tests for loading text from supported healthcare document formats."""

import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from docx import Document

from mediassist_healthcare.ingestion import load_document


class LoadDocumentTests(unittest.TestCase):
    """Check that the loader handles supported and invalid file inputs."""

    # Check that a TXT file returns its text and source metadata.
    def test_loads_utf8_text_and_source_metadata(self) -> None:
        """Verify text files return their content and source details."""

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "policy.txt"
            path.write_text("Admission policy", encoding="utf-8")

            result = load_document(path)

        self.assertEqual(result.source, path)
        self.assertEqual(result.file_type, "txt")
        self.assertEqual(result.text, "Admission policy")

    # Check that the DOCX loader reads paragraphs and table content.
    def test_loads_docx_paragraphs_and_table_content(self) -> None:
        """Verify Word paragraphs and table cells are both extracted."""

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "report.docx"
            document = Document()
            document.add_paragraph("Patient summary")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Test"
            table.cell(0, 1).text = "Result"
            document.save(path)

            result = load_document(path)

        self.assertEqual(result.file_type, "docx")
        self.assertIn("Patient summary", result.text)
        self.assertIn("Test\tResult", result.text)

    @patch("mediassist_healthcare.ingestion.PdfReader")
    # Check that PDF pages are extracted and kept in order.
    def test_loads_pdf_page_text(self, pdf_reader: Mock) -> None:
        """Verify extracted PDF page text is joined in page order."""

        first_page = Mock()
        first_page.extract_text.return_value = "Page one"
        second_page = Mock()
        second_page.extract_text.return_value = "Page two"
        pdf_reader.return_value.pages = [first_page, second_page]

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "report.pdf"
            path.touch()

            result = load_document(path)

        self.assertEqual(result.file_type, "pdf")
        self.assertEqual(result.text, "Page one\nPage two")

    # Check that unsupported file formats are rejected with a clear message.
    def test_rejects_unsupported_file_type(self) -> None:
        """Verify an unsupported file extension produces a clear error."""

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "notes.csv"
            path.write_text("not a supported document", encoding="utf-8")

            with self.assertRaisesRegex(ValueError, "Supported types are PDF, DOCX, and TXT"):
                load_document(path)

    # Check that requesting a missing file raises a file-not-found error.
    def test_rejects_missing_file(self) -> None:
        """Verify a missing input path is not silently accepted."""

        with self.assertRaises(FileNotFoundError):
            load_document(Path("missing-document.txt"))


# Allow this test file to run directly as well as through test discovery.
if __name__ == "__main__":
    unittest.main()
